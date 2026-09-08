from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree

from conversion_engine.builders.docx import DocxBuilder
from conversion_engine.domain.models import (
    BBox,
    DocumentIR,
    PageIR,
    ParagraphBlock,
    TableBlock,
    TextSpan,
)


def paragraph(text: str, kind: str = "paragraph") -> ParagraphBlock:
    return ParagraphBlock(
        spans=[TextSpan(text=text, bbox=BBox(36, 36, 300, 50), bold=kind == "heading")],
        bbox=BBox(36, 36, 300, 50),
        kind=kind,
        heading_level=1 if kind == "heading" else None,
        list_ordered=kind == "list",
    )


def test_builds_editable_paragraphs_table_and_page_break(tmp_path: Path):
    target = tmp_path / "result.docx"
    source = DocumentIR(
        pages=[
            PageIR(
                number=1,
                width=595,
                height=842,
                blocks=[
                    paragraph("Editable heading", "heading"),
                    paragraph("1. Editable list", "list"),
                    TableBlock([["Item", "Price"], ["A", "10"]], BBox(36, 100, 400, 180)),
                ],
            ),
            PageIR(
                number=2,
                width=595,
                height=842,
                blocks=[paragraph("Second page")],
            ),
        ]
    )
    DocxBuilder().build(source, target)

    output = Document(target)
    with ZipFile(target) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    all_text = " ".join(
        root.xpath(
            ".//w:t/text()",
            namespaces={
                "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            },
        )
    )
    assert "Editable heading" in all_text
    assert "Editable list" in all_text
    assert output.tables[0].cell(1, 1).text == "10"
    assert len(output.sections) == 2


def test_dense_pages_keep_source_positions_and_footer(tmp_path: Path):
    target = tmp_path / "positioned.docx"

    def positioned(text: str, top: float) -> ParagraphBlock:
        bbox = BBox(42, top, 550, top + 14)
        return ParagraphBlock(
            spans=[TextSpan(text=text, bbox=bbox, font_size=10)],
            bbox=bbox,
        )

    source = DocumentIR(
        pages=[
            PageIR(
                number=number,
                width=595,
                height=842,
                blocks=[
                    positioned(f"Chapter {number}", 60),
                    positioned("First paragraph", 110),
                    positioned("Second paragraph", 150),
                    positioned(f"Report | Page {number}", 796),
                ],
            )
            for number in (1, 2)
        ]
    )
    DocxBuilder().build(source, target)

    with ZipFile(target) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    namespaces = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "v": "urn:schemas-microsoft-com:vml",
    }
    textboxes = root.xpath(".//v:textbox", namespaces=namespaces)
    footer_shapes = root.xpath(
        './/v:shape[contains(@style, "margin-top:796.00pt")]',
        namespaces=namespaces,
    )

    assert len(Document(target).sections) == 2
    assert len(textboxes) == 8
    assert len(footer_shapes) == 2


def test_normalizes_pdf_font_face_names():
    assert DocxBuilder._safe_font("ABCDEF+Arial-Bold") == "Arial"
    assert DocxBuilder._safe_font("Carlito-Regular") == "Arial"
    assert DocxBuilder._word_font_size("Carlito-Regular", 11) == 9.9
