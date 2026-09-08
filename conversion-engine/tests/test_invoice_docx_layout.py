from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from conversion_engine.infra.logging import configure_logging
from conversion_engine.services.conversion import ConversionRequest, ConversionService


def _invoice_pdf(path: Path) -> None:
    pdf = canvas.Canvas(str(path), pagesize=A4)
    pdf.setFont("Helvetica-Bold", 20)
    pdf.drawRightString(560, 805, "DEVIS")
    pdf.setFont("Helvetica-Bold", 9)
    pdf.drawString(34, 700, "ÉMETTEUR")
    pdf.drawString(312, 700, "DESTINATAIRE")
    pdf.drawString(34, 684, "ABC SARL")
    pdf.drawString(312, 684, "SOPAL")
    pdf.setFont("Helvetica", 9)
    pdf.drawString(34, 668, "Espace Tunis immeuble H. Bureau B3-1")
    pdf.drawString(312, 668, "Route de Gabes KM 1.5 3003 Sfax")
    pdf.drawString(34, 652, "Tél: +216 55 053 505")
    pdf.drawString(312, 652, "contact@sopal.example")

    headers = [
        ("Description", 31),
        ("Prix unit.", 270),
        ("Qté", 330),
        ("Remise", 378),
        ("Total HT", 472),
        ("TVA", 527),
    ]
    pdf.setFont("Helvetica-Bold", 9)
    for value, x in headers:
        pdf.drawString(x, 610, value)
    pdf.setFont("Helvetica", 9)
    values = [
        ("Page publicitaire intérieure préférentielle", 29),
        ("3000.000 DT", 257),
        ("1", 335),
        ("10%", 386),
        ("2700.000 DT", 457),
        ("19%", 529),
    ]
    for value, x in values:
        pdf.drawString(x, 590, value)
    pdf.save()


def test_reconstructs_borderless_invoice_as_editable_layout(tmp_path: Path):
    source = tmp_path / "invoice.pdf"
    target = tmp_path / "invoice.docx"
    _invoice_pdf(source)

    result = ConversionService(configure_logging()).convert(
        ConversionRequest(source, target)
    )
    output = Document(target)

    with ZipFile(target) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    namespaces = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "v": "urn:schemas-microsoft-com:vml",
    }
    textboxes = [
        " ".join(box.xpath(".//w:t/text()", namespaces=namespaces))
        for box in root.xpath(".//v:textbox", namespaces=namespaces)
    ]
    assert any(
        "ÉMETTEUR" in value and "ABC SARL" in value for value in textboxes
    )
    assert any(
        "DESTINATAIRE" in value and "SOPAL" in value for value in textboxes
    )

    item_table = next(table for table in output.tables if len(table.columns) == 6)
    assert [cell.text for cell in item_table.rows[0].cells] == [
        "Description",
        "Prix unit.",
        "Qté",
        "Remise",
        "Total HT",
        "TVA",
    ]
    assert [cell.text for cell in item_table.rows[1].cells][1:] == [
        "3000.000 DT",
        "1",
        "10%",
        "2700.000 DT",
        "19%",
    ]
    assert result.report.tables == 1
    assert result.report.text_coverage >= 0.98
