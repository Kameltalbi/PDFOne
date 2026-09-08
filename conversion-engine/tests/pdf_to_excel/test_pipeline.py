from __future__ import annotations

from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from conversion_engine.domain.errors import InvalidPdfError
from conversion_engine.infra.logging import configure_logging
from conversion_engine.pdf_to_excel.converter import PdfToExcelConverter
from conversion_engine.pdf_to_excel.excel_builder import ExcelBuilder, ValueParser
from conversion_engine.pdf_to_excel.models import ConversionRequest


def convert(source: Path, tmp_path: Path):
    output = tmp_path / f"{source.stem}.xlsx"
    report = tmp_path / f"{source.stem}.quality.json"
    result = PdfToExcelConverter(configure_logging()).convert(
        ConversionRequest(source, output, report)
    )
    return result, load_workbook(output, data_only=False)


def table_sheet(workbook):
    return workbook[workbook.sheetnames[1]]


def all_values(sheet):
    return [
        cell.value
        for row in sheet.iter_rows()
        for cell in row
        if cell.value not in (None, "")
    ]


def test_softfacture_invoice_regression(pdf_fixtures, tmp_path):
    result, workbook = convert(pdf_fixtures["softfacture"], tmp_path)
    sheet = table_sheet(workbook)
    assert [cell.value for cell in sheet[1]] == [
        "Description", "Prix unit.", "Qté", "Remise", "Total HT", "TVA"
    ]
    assert sheet["A2"].value == "Page publicité 1"
    assert sheet["B2"].value == pytest.approx(3000.0)
    assert sheet["C2"].value == 1
    assert sheet["D2"].value == pytest.approx(0.10)
    assert sheet["D2"].number_format == "0.00%"
    assert sheet["E2"].value == pytest.approx(2700.0)
    assert sheet["F2"].value == pytest.approx(0.19)
    assert sheet["A1"].font.bold
    overview = all_values(workbook["Document"])
    assert "SoftFacture SARL" in overview
    assert "ACME Customer" in overview
    assert any("SF-2026-001" in str(value) for value in overview)
    assert any("Total TTC" in str(value) for value in overview)
    assert result.diagnostics.tables_detected == 1
    assert result.diagnostics.numeric_preservation == 1.0
    assert result.diagnostics.text_coverage >= 0.95


def test_quotation_preserves_metadata_and_parties(pdf_fixtures, tmp_path):
    result, workbook = convert(pdf_fixtures["quotation"], tmp_path)
    values = all_values(workbook["Document"])
    assert any("DEVIS" in str(value) for value in values)
    assert "SoftFacture SARL" in values
    assert "ACME Customer" in values
    assert result.diagnostics.conversion_quality >= 0.85


def test_bordered_table(pdf_fixtures, tmp_path):
    result, workbook = convert(pdf_fixtures["bordered"], tmp_path)
    sheet = table_sheet(workbook)
    assert [cell.value for cell in sheet[1]] == ["Item", "Qty", "Amount"]
    assert sheet["B2"].value == 2
    assert sheet["C3"].value == pytest.approx(45.0)
    assert "bordered" in result.diagnostics.strategies


def test_borderless_table(pdf_fixtures, tmp_path):
    result, workbook = convert(pdf_fixtures["borderless"], tmp_path)
    sheet = table_sheet(workbook)
    assert sheet.max_column == 3
    assert sheet.max_row == 4
    assert sheet["B3"].value == 4
    assert "semantic-alignment" in result.diagnostics.strategies


def test_multi_page_repeated_headers_are_merged(pdf_fixtures, tmp_path):
    result, workbook = convert(pdf_fixtures["multipage"], tmp_path)
    assert result.diagnostics.tables_detected == 1
    assert len(workbook.sheetnames) == 2
    sheet = table_sheet(workbook)
    assert sheet.max_row == 5
    assert sum(cell.value == "Description" for cell in sheet["A"]) == 1


def test_several_independent_tables(pdf_fixtures, tmp_path):
    result, workbook = convert(pdf_fixtures["multiple"], tmp_path)
    assert result.diagnostics.tables_detected == 2
    assert len(workbook.sheetnames) == 3
    assert workbook[workbook.sheetnames[2]]["A1"].value == "Tax"


def test_landscape_pdf(pdf_fixtures, tmp_path):
    result, workbook = convert(pdf_fixtures["landscape"], tmp_path)
    assert result.diagnostics.pages == 1
    assert table_sheet(workbook).max_column == 3


def test_pdf_with_text_and_image(pdf_fixtures, tmp_path):
    result, workbook = convert(pdf_fixtures["image"], tmp_path)
    assert result.diagnostics.document_type == "mixed"
    assert table_sheet(workbook)["A2"].value == "Visual service"


def test_malformed_pdf_is_rejected(pdf_fixtures, tmp_path):
    with pytest.raises(InvalidPdfError):
        convert(pdf_fixtures["malformed"], tmp_path)


def test_value_types_are_conservative():
    parser = ValueParser()
    assert parser.parse("3000.000")[0] == pytest.approx(3000.0)
    assert parser.parse("10%") == (pytest.approx(0.1), "0.00%")
    assert str(parser.parse("17/09/2026")[0]) == "2026-09-17"
    assert parser.parse("07/09/2026") == ("07/09/2026", None)


def test_formula_like_source_text_is_not_executable():
    cell = Workbook().active["A1"]
    ExcelBuilder._set_cell(cell, "=HYPERLINK(\"https://example.test\")", None)
    assert cell.value == '=HYPERLINK("https://example.test")'
    assert cell.data_type == "s"
    assert cell.quotePrefix


def test_table_does_not_duplicate_worksheet_auto_filter(pdf_fixtures, tmp_path):
    result, _workbook = convert(pdf_fixtures["softfacture"], tmp_path)
    namespace = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with ZipFile(result.output_path) as archive:
        worksheet = ElementTree.fromstring(archive.read("xl/worksheets/sheet2.xml"))
        table = ElementTree.fromstring(archive.read("xl/tables/table1.xml"))
    assert worksheet.find("x:autoFilter", namespace) is None
    assert table.find("x:autoFilter", namespace) is not None
