from __future__ import annotations

import json
import re
from dataclasses import asdict, dataclass
from difflib import SequenceMatcher
from pathlib import Path

from docx import Document

from conversion_engine.domain.models import (
    DocumentIR,
    ImageBlock,
    ParagraphBlock,
    TableBlock,
)


def _normalize(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().casefold()


@dataclass
class QualityReport:
    text_similarity: float
    text_coverage: float
    source_characters: int
    output_characters: int
    pages: int
    paragraphs: int
    headings: int
    lists: int
    tables: int
    images: int
    warnings: list[str]

    def as_dict(self) -> dict:
        return asdict(self)


class QualityEvaluator:
    def evaluate(self, source: DocumentIR, docx_path: Path) -> QualityReport:
        document = Document(docx_path)
        output_parts = [paragraph.text for paragraph in document.paragraphs]
        output_parts.extend(
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        )
        original = _normalize(source.text)
        output = _normalize("\n".join(output_parts))
        paragraph_blocks = [
            block
            for page in source.pages
            for block in page.blocks
            if isinstance(block, ParagraphBlock)
        ]
        return QualityReport(
            text_similarity=round(SequenceMatcher(None, original, output).ratio(), 4)
            if original or output
            else 1.0,
            text_coverage=round(min(1.0, len(output) / max(1, len(original))), 4),
            source_characters=len(original),
            output_characters=len(output),
            pages=len(source.pages),
            paragraphs=len(paragraph_blocks),
            headings=sum(block.kind == "heading" for block in paragraph_blocks),
            lists=sum(block.kind == "list" for block in paragraph_blocks),
            tables=sum(
                isinstance(block, TableBlock)
                for page in source.pages
                for block in page.blocks
            ),
            images=sum(
                isinstance(block, ImageBlock)
                for page in source.pages
                for block in page.blocks
            ),
            warnings=source.warnings,
        )

    @staticmethod
    def write(report: QualityReport, target: Path) -> None:
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(
            json.dumps(report.as_dict(), ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
