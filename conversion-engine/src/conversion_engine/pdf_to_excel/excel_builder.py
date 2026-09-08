from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Any, List, Optional, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table as ExcelTable
from openpyxl.worksheet.table import TableStyleInfo

from .models import CellValue, DocumentAnalysis, SemanticZone, Table


CURRENCY_RE = re.compile(
    r"^\s*(?:(EUR|USD|GBP|TND|CAD|MAD|CHF)|([€$£]))?\s*"
    r"([+-]?(?:\d{1,3}(?:[ .,'’]\d{3})+|\d+)(?:[.,]\d+)?)"
    r"\s*(?:(EUR|USD|GBP|TND|CAD|MAD|CHF)|([€$£]))?\s*$",
    re.I,
)
PERCENT_RE = re.compile(r"^\s*([+-]?\d+(?:[.,]\d+)?)\s*%\s*$")
DATE_RE = re.compile(r"^\s*(\d{1,2})/(\d{1,2})/(\d{4})\s*$")


class ValueParser:
    CURRENCY_SYMBOLS = {"€": "EUR", "$": "USD", "£": "GBP"}

    def parse(self, raw: str) -> CellValue:
        value = raw.strip()
        percent = PERCENT_RE.match(value)
        if percent:
            number = float(percent.group(1).replace(",", ".")) / 100
            return number, "0.00%"
        date = DATE_RE.match(value)
        if date:
            day, month, year = map(int, date.groups())
            if day > 12 or month > 12:
                try:
                    return datetime(year, month, day).date(), "dd/mm/yyyy"
                except ValueError:
                    return raw, None
        currency = CURRENCY_RE.match(value)
        if currency:
            code = currency.group(1) or currency.group(4)
            symbol = currency.group(2) or currency.group(5)
            parsed = self._number(currency.group(3))
            if parsed is not None:
                if code or symbol:
                    currency_code = (code or self.CURRENCY_SYMBOLS.get(symbol, "")).upper()
                    return parsed, f'#,##0.00 "{"€" if currency_code == "EUR" else currency_code}"'
                return parsed, "#,##0.###"
        return raw, None

    @staticmethod
    def _number(value: str) -> Optional[Any]:
        compact = value.replace(" ", "").replace("'", "").replace("’", "")
        if "," in compact and "." in compact:
            decimal = "," if compact.rfind(",") > compact.rfind(".") else "."
            thousands = "." if decimal == "," else ","
            compact = compact.replace(thousands, "").replace(decimal, ".")
        elif "," in compact:
            compact = compact.replace(",", ".")
        try:
            number = float(compact)
            return int(number) if number.is_integer() and "." not in compact else number
        except ValueError:
            return None


class ExcelBuilder:
    HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
    SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
    WHITE_FONT = Font(color="FFFFFF", bold=True)
    THIN = Side(style="thin", color="B7C9D6")

    def __init__(self):
        self.values = ValueParser()

    def build(self, analysis: DocumentAnalysis, tables: Sequence[Table], output: Path) -> None:
        workbook = Workbook()
        overview = workbook.active
        overview.title = "Document"
        self._build_overview(overview, analysis.zones)
        for index, table in enumerate(tables, start=1):
            sheet = workbook.create_sheet(self._sheet_name(table, index))
            self._build_table_sheet(sheet, table, index)
        if not tables:
            sheet = workbook.create_sheet("Extracted content")
            sheet.append(["Page", "Content"])
            for zone in analysis.zones:
                for line in zone.lines:
                    sheet.append([zone.page_number, line])
                    self._set_cell(sheet.cell(sheet.max_row, 2), line, None)
            self._style_header(sheet, 1, 2)
            sheet.column_dimensions["A"].width = 10
            sheet.column_dimensions["B"].width = 90
            sheet.freeze_panes = "A2"
        output.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output)

    def _build_overview(self, sheet, zones: Sequence[SemanticZone]) -> None:
        sheet.sheet_view.showGridLines = False
        sheet.merge_cells("A1:G1")
        sheet["A1"] = "Converted document"
        sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
        sheet["A1"].fill = self.HEADER_FILL
        sheet["A1"].alignment = Alignment(horizontal="center")
        row = 3
        grouped = {kind: [zone for zone in zones if zone.kind == kind] for kind in (
            "document_header", "seller", "customer", "totals", "other", "footer"
        )}
        row = self._zone_section(sheet, row, "Document information", grouped["document_header"], 1, 7)

        seller_lines = self._zone_lines(grouped["seller"])
        customer_lines = self._zone_lines(grouped["customer"])
        if seller_lines or customer_lines:
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
            sheet.merge_cells(start_row=row, start_column=5, end_row=row, end_column=7)
            sheet.cell(row, 1, "Seller").fill = self.SECTION_FILL
            sheet.cell(row, 5, "Customer").fill = self.SECTION_FILL
            sheet.cell(row, 1).font = sheet.cell(row, 5).font = Font(bold=True)
            row += 1
            height = max(len(seller_lines), len(customer_lines), 1)
            for offset in range(height):
                if offset < len(seller_lines):
                    sheet.merge_cells(start_row=row + offset, start_column=1, end_row=row + offset, end_column=3)
                    self._set_cell(sheet.cell(row + offset, 1), seller_lines[offset], None)
                if offset < len(customer_lines):
                    sheet.merge_cells(start_row=row + offset, start_column=5, end_row=row + offset, end_column=7)
                    self._set_cell(sheet.cell(row + offset, 5), customer_lines[offset], None)
            row += height + 1

        row = self._zone_section(sheet, row, "Totals", grouped["totals"], 4, 7, typed=True)
        row = self._zone_section(sheet, row, "Additional content", grouped["other"], 1, 7)
        self._zone_section(sheet, row, "Footer / legal information", grouped["footer"], 1, 7)
        for column in range(1, 8):
            sheet.column_dimensions[get_column_letter(column)].width = 18
        for cells in sheet.iter_rows():
            for cell in cells:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    def _zone_section(
        self, sheet, row: int, title: str, zones: Sequence[SemanticZone],
        start_column: int, end_column: int, typed: bool = False
    ) -> int:
        lines = self._zone_lines(zones)
        if not lines:
            return row
        sheet.merge_cells(
            start_row=row, start_column=start_column, end_row=row, end_column=end_column
        )
        heading = sheet.cell(row, start_column, title)
        heading.fill = self.SECTION_FILL
        heading.font = Font(bold=True)
        row += 1
        for line in lines:
            if typed:
                label, value = self._split_label_value(line)
                sheet.cell(row, max(start_column, end_column - 2), label)
                cell = sheet.cell(row, end_column)
                parsed, number_format = self.values.parse(value)
                self._set_cell(cell, parsed, number_format)
            else:
                sheet.merge_cells(
                    start_row=row,
                    start_column=start_column,
                    end_row=row,
                    end_column=end_column,
                )
                self._set_cell(sheet.cell(row, start_column), line, None)
            row += 1
        return row + 1

    def _build_table_sheet(self, sheet, table: Table, index: int) -> None:
        sheet.sheet_view.showGridLines = False
        headers = table.headers or [f"Column {column + 1}" for column in range(max(map(len, table.rows), default=1))]
        sheet.append(headers)
        for column, header in enumerate(headers, start=1):
            self._set_cell(sheet.cell(1, column), header, None)
        for row in table.rows:
            values: List[Any] = []
            formats: List[Optional[str]] = []
            for column in range(len(headers)):
                raw = row[column] if column < len(row) else ""
                value, number_format = self.values.parse(raw)
                values.append(value)
                formats.append(number_format)
            sheet.append(values)
            for column, number_format in enumerate(formats, start=1):
                self._set_cell(
                    sheet.cell(sheet.max_row, column),
                    values[column - 1],
                    number_format,
                )
        self._style_header(sheet, 1, len(headers))
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.border = Border(bottom=self.THIN)
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                    horizontal="right" if isinstance(cell.value, (int, float)) else "left",
                )
        sheet.freeze_panes = "A2"
        if sheet.max_row >= 2:
            excel_table = ExcelTable(
                displayName=f"ExtractedTable{index}",
                ref=f"A1:{get_column_letter(len(headers))}{sheet.max_row}",
            )
            excel_table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            sheet.add_table(excel_table)
        else:
            # A worksheet filter and a table filter over the same range make
            # some Microsoft Excel versions report a corrupted workbook.
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
        for column, header in enumerate(headers, start=1):
            values = [str(sheet.cell(row, column).value or "") for row in range(1, sheet.max_row + 1)]
            width = min(60, max(10, max(map(len, values), default=len(header)) + 2))
            sheet.column_dimensions[get_column_letter(column)].width = width

    def _style_header(self, sheet, row: int, columns: int) -> None:
        for column in range(1, columns + 1):
            cell: Cell = sheet.cell(row, column)
            cell.fill = self.HEADER_FILL
            cell.font = self.WHITE_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=self.THIN)
        sheet.row_dimensions[row].height = 30

    @staticmethod
    def _set_cell(cell: Cell, value: Any, number_format: Optional[str]) -> None:
        cell.value = value
        if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
            cell.data_type = "s"
            cell.quotePrefix = True
        if number_format:
            cell.number_format = number_format

    @staticmethod
    def _zone_lines(zones: Sequence[SemanticZone]) -> List[str]:
        return [line for zone in zones for line in zone.lines]

    @staticmethod
    def _split_label_value(line: str) -> Tuple[str, str]:
        match = re.match(r"^\s*(.+?)(?:\s*[:：]\s*|\s{2,})([-+€$£\d].*)$", line)
        if match:
            return match.group(1).strip(), match.group(2).strip()
        tokens = line.rsplit(" ", 1)
        return (tokens[0], tokens[1]) if len(tokens) == 2 else (line, "")

    @staticmethod
    def _sheet_name(table: Table, index: int) -> str:
        base = "Items" if index == 1 else f"Table {index}"
        if len(table.source_pages) > 1:
            base += f" p{table.source_pages[0]}-{table.source_pages[-1]}"
        return base[:31]
