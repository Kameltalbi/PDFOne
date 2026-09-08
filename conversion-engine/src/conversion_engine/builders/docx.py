from __future__ import annotations

import re
from statistics import median
from pathlib import Path
from typing import Iterable, List, Sequence, Tuple

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree

from conversion_engine.domain.models import (
    Block,
    DocumentIR,
    ImageBlock,
    PageIR,
    ParagraphBlock,
    TableBlock,
)


LIST_PREFIX = re.compile(r"^\s*(?:(?:\d+|[A-Za-z])[.)]|[•●▪◦‣–-])\s+")
VML_NS = "urn:schemas-microsoft-com:vml"
OFFICE_NS = "urn:schemas-microsoft-com:office:office"


class DocxBuilder:
    def build(self, source: DocumentIR, output: Path) -> None:
        document = Document()
        self._shape_id = 1
        self._vml_shape_type_added = False
        if source.title:
            document.core_properties.title = source.title
        if source.author:
            document.core_properties.author = source.author
        document.core_properties.subject = "Converted by One2PDF"
        self._configure_styles(document)

        for index, page in enumerate(source.pages):
            section = (
                document.sections[0]
                if index == 0
                else document.add_section(WD_SECTION.NEW_PAGE)
            )
            if index:
                separator = document.paragraphs[-1]
                separator.paragraph_format.space_before = Pt(0)
                separator.paragraph_format.space_after = Pt(0)
                separator.paragraph_format.line_spacing = Pt(1)
                separator.paragraph_format.keep_with_next = False
                separator.add_run().font.size = Pt(1)
            body_blocks, footer_blocks = self._split_footer_blocks(page)
            self._configure_section(section, page, body_blocks)
            self._add_section_footer(section, footer_blocks, page)
            if self._use_fixed_layout(page):
                self._add_fixed_page(document, page)
                continue
            # Flowing Word paragraphs for reports: Microsoft Word clips
            # absolute VML textboxes on the right edge, so dense text pages
            # must use real paragraphs that wrap inside the section margins.
            previous_bottom = float(section.top_margin.pt)
            for band in self._layout_bands(body_blocks):
                top = min(block.bbox.top for block in band)
                gap = max(0.0, top - previous_bottom - 4.0)
                columns = self._layout_columns(band)
                if len(columns) > 1:
                    self._add_vertical_gap(document, gap)
                    self._add_layout_row(document, columns, page)
                else:
                    for block in sorted(columns[0], key=lambda item: item.bbox.top):
                        self._add_block(document, block, page)
                        if document.paragraphs and not isinstance(block, TableBlock):
                            document.paragraphs[-1].paragraph_format.space_before = Pt(gap)
                        gap = 0.0
                previous_bottom = max(block.bbox.bottom for block in band)

        output.parent.mkdir(parents=True, exist_ok=True)
        document.save(output)

    @staticmethod
    def _configure_styles(document: Document) -> None:
        normal = document.styles["Normal"]
        normal.font.name = "Arial"
        normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(5)
        for level in range(1, 4):
            style = document.styles[f"Heading {level}"]
            style.font.name = "Arial"
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.font.size = Pt({1: 20, 2: 16, 3: 13}[level])
            style.paragraph_format.keep_with_next = True

    @staticmethod
    def _configure_section(section, page: PageIR, body_blocks: Sequence[Block] = ()) -> None:
        section.page_width = Pt(page.width)
        section.page_height = Pt(page.height)
        content = [
            block
            for block in body_blocks
            if isinstance(block, (ParagraphBlock, TableBlock, ImageBlock))
        ]
        if content:
            left = min(block.bbox.x0 for block in content)
            right = page.width - max(block.bbox.x1 for block in content)
            top = min(block.bbox.top for block in content)
            bottom = page.height - max(block.bbox.bottom for block in content)
        else:
            left = right = top = bottom = 24.0
        section.top_margin = Pt(max(18.0, min(96.0, top)))
        section.bottom_margin = Pt(max(28.0, min(96.0, bottom - 12 if bottom > 18 else 48.0)))
        section.left_margin = Pt(max(24.0, min(90.0, left)))
        section.right_margin = Pt(max(24.0, min(90.0, right)))

    @staticmethod
    def _use_fixed_layout(page: PageIR) -> bool:
        paragraphs = [
            block for block in page.blocks if isinstance(block, ParagraphBlock)
        ]
        tables = any(isinstance(block, TableBlock) for block in page.blocks)
        if len(paragraphs) < 3:
            return False
        # Long body copy must stay in flowing Word paragraphs. Absolute VML
        # textboxes are clipped on the right by Microsoft Word even when the
        # stored width looks correct.
        long_count = sum(
            1
            for block in paragraphs
            if block.bbox.width >= page.width * 0.65
            and (
                len(block.text) > 100
                or block.bbox.height >= max(
                    28.0,
                    max((span.font_size for span in block.spans), default=11) * 2.2,
                )
            )
        )
        if long_count >= 2:
            return False
        # Invoices keep absolute label placement around real Word tables.
        if tables:
            return True
        return long_count <= 1

    @staticmethod
    def _split_footer_blocks(page: PageIR) -> Tuple[List[Block], List[Block]]:
        body: List[Block] = []
        footer: List[Block] = []
        for block in page.blocks:
            if DocxBuilder._is_footer_block(block, page):
                footer.append(block)
            else:
                body.append(block)
        return body, footer

    @staticmethod
    def _is_footer_block(block: Block, page: PageIR) -> bool:
        if block.bbox.top < page.height * 0.88:
            return False
        if isinstance(block, ImageBlock):
            return block.bbox.height <= 40
        if isinstance(block, ParagraphBlock):
            return bool(re.search(r"(?:\bpage\s+\d+|^\s*\d+\s*$)", block.text, re.I))
        return False

    def _add_section_footer(self, section, footer_blocks: Sequence[Block], page: PageIR) -> None:
        footer = section.footer
        footer.is_linked_to_previous = False
        paragraph = footer.paragraphs[0]
        paragraph.style = "Normal"
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        if not footer_blocks:
            return
        section.footer_distance = Pt(max(12, page.height - max(b.bbox.bottom for b in footer_blocks)))
        for index, band in enumerate(self._layout_bands(footer_blocks)):
            if index:
                paragraph = footer.add_paragraph(style="Normal")
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
            for block in sorted(band, key=lambda item: item.bbox.x0):
                offset = max(0, block.bbox.x0 - section.left_margin.pt)
                if offset:
                    paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(offset))
                    paragraph.add_run("\t")
                if isinstance(block, ParagraphBlock):
                    for span in block.spans:
                        run = paragraph.add_run(span.text.replace("\n", " "))
                        run.bold = span.bold
                        run.italic = span.italic
                        run.font.name = self._safe_font(span.font_name)
                        run.font.size = Pt(self._word_font_size(span.font_name, span.font_size))
                        if span.color:
                            run.font.color.rgb = RGBColor(*span.color)
                elif isinstance(block, ImageBlock) and block.path.exists():
                    paragraph.add_run().add_picture(
                        str(block.path), width=Pt(block.bbox.width), height=Pt(block.bbox.height),
                    )

    def _add_fixed_page(self, document: Document, page: PageIR) -> None:
        """Build business documents on a page canvas with editable objects."""
        body_blocks, _footer_blocks = self._split_footer_blocks(page)
        anchor = document.add_paragraph()
        anchor.paragraph_format.space_before = Pt(0)
        first_table_top = min(
            (
                block.bbox.top
                for block in body_blocks
                if isinstance(block, TableBlock)
            ),
            default=24.0,
        )
        anchor.paragraph_format.space_after = Pt(max(0, first_table_top - document.sections[-1].top_margin.pt - 2))
        anchor.paragraph_format.line_spacing = Pt(1)
        anchor.add_run(" ").font.size = Pt(1)

        for block in body_blocks:
            if isinstance(block, ParagraphBlock):
                self._add_floating_textbox(anchor, block, page)
            elif isinstance(block, ImageBlock):
                self._add_floating_image(anchor, block)
            elif isinstance(block, TableBlock):
                self._add_table(document, block, document.sections[-1])

    def _add_floating_textbox(
        self,
        anchor,
        block: ParagraphBlock,
        page: PageIR,
    ) -> None:
        run = anchor.add_run()
        pict = OxmlElement("w:pict")
        if not self._vml_shape_type_added:
            shape_type = etree.Element(f"{{{VML_NS}}}shapetype")
            shape_type.set("id", "_x0000_t202")
            shape_type.set("coordsize", "21600,21600")
            shape_type.set(f"{{{OFFICE_NS}}}spt", "202")
            shape_type.set("path", "m,l,21600r21600,l21600,xe")
            stroke = etree.Element(f"{{{VML_NS}}}stroke")
            stroke.set("joinstyle", "miter")
            shape_type.append(stroke)
            path = etree.Element(f"{{{VML_NS}}}path")
            path.set("gradientshapeok", "t")
            path.set(f"{{{OFFICE_NS}}}connecttype", "rect")
            shape_type.append(path)
            pict.append(shape_type)
            self._vml_shape_type_added = True

        font_size = max((span.font_size for span in block.spans), default=9)
        # Keep the PDF's right margin. A large safety width makes VML text
        # boxes run to the page edge in Microsoft Word even when LibreOffice
        # renders them correctly.
        width_padding = min(3.0, max(1.5, (page.width - block.bbox.x1) * 0.05))
        if block.kind == "heading" and block.alignment == "left":
            width_padding = max(width_padding, font_size * 3.0)
        width = min(
            max(14.0, block.bbox.width + width_padding),
            max(14.0, page.width - block.bbox.x0 - 24.0),
        )
        # Centered cover labels need symmetric room for font substitution.
        left = block.bbox.x0
        if block.alignment == "center":
            center = (block.bbox.x0 + block.bbox.x1) / 2
            width = max(width, 2 * min(center - 24, page.width - 24 - center))
            left = center - width / 2
        height = max(font_size * 1.35, block.bbox.height + 5)
        shape = etree.Element(f"{{{VML_NS}}}shape")
        shape.set("id", f"_x0000_s{1024 + self._shape_id}")
        self._shape_id += 1
        shape.set("type", "#_x0000_t202")
        shape.set(
            "style",
            (
                f"position:absolute;margin-left:{left:.2f}pt;"
                f"margin-top:{block.bbox.top:.2f}pt;width:{width:.2f}pt;"
                f"height:{height:.2f}pt;mso-position-horizontal-relative:page;"
                "mso-position-vertical-relative:page;z-index:251658240;"
                "mso-wrap-style:none;mso-wrap-distance-left:0;"
                "mso-wrap-distance-right:0"
            ),
        )
        shape.set("stroked", "f")
        shape.set("filled", "f")

        textbox = etree.Element(f"{{{VML_NS}}}textbox")
        textbox.set("inset", "0,0,0,0")
        content = OxmlElement("w:txbxContent")
        paragraph = OxmlElement("w:p")
        properties = OxmlElement("w:pPr")
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:before"), "0")
        spacing.set(qn("w:after"), "0")
        line_count = max(
            1,
            1 + sum(span.text.count("\n") for span in block.spans),
        )
        source_line_height = (
            max(font_size * 1.05, (block.bbox.height - font_size) / (line_count - 1))
            if line_count > 1
            else font_size * 1.15
        )
        spacing.set(qn("w:line"), str(round(source_line_height * 20)))
        spacing.set(qn("w:lineRule"), "exact")
        properties.append(spacing)
        if block.alignment != "left":
            alignment = OxmlElement("w:jc")
            alignment.set(
                qn("w:val"),
                {"center": "center", "right": "right", "justify": "both"}.get(
                    block.alignment, "left"
                ),
            )
            properties.append(alignment)
        paragraph.append(properties)
        for span in block.spans:
            self._append_vml_text_run(paragraph, span)
        content.append(paragraph)
        textbox.append(content)
        shape.append(textbox)
        pict.append(shape)
        run._r.append(pict)

    @staticmethod
    def _append_vml_text_run(paragraph, span) -> None:
        run = OxmlElement("w:r")
        properties = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        font_name = DocxBuilder._safe_font(span.font_name)
        fonts.set(qn("w:ascii"), font_name)
        fonts.set(qn("w:hAnsi"), font_name)
        properties.append(fonts)
        size = OxmlElement("w:sz")
        word_size = DocxBuilder._word_font_size(span.font_name, span.font_size)
        size.set(qn("w:val"), str(round(word_size * 2)))
        properties.append(size)
        if span.bold:
            properties.append(OxmlElement("w:b"))
        if span.italic:
            properties.append(OxmlElement("w:i"))
        if span.color and not all(component > 245 for component in span.color):
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "".join(f"{component:02X}" for component in span.color))
            properties.append(color)
        run.append(properties)
        pieces = span.text.split("\n")
        for index, piece in enumerate(pieces):
            if index:
                run.append(OxmlElement("w:br"))
            if piece:
                text = OxmlElement("w:t")
                text.set(qn("xml:space"), "preserve")
                text.text = piece
                run.append(text)
        paragraph.append(run)

    def _add_floating_image(self, anchor, block: ImageBlock) -> None:
        if not block.path.exists():
            return
        run = anchor.add_run()
        run.add_picture(
            str(block.path),
            width=Pt(max(12, block.bbox.width)),
            height=Pt(max(12, block.bbox.height)),
        )
        inline = run._r.find(".//" + qn("wp:inline"))
        if inline is None:
            return
        floating = OxmlElement("wp:anchor")
        for name, value in {
            "distT": "0",
            "distB": "0",
            "distL": "0",
            "distR": "0",
            "simplePos": "0",
            "relativeHeight": str(251658240 + self._shape_id),
            "behindDoc": "0",
            "locked": "0",
            "layoutInCell": "1",
            "allowOverlap": "1",
        }.items():
            floating.set(name, value)
        self._shape_id += 1
        simple = OxmlElement("wp:simplePos")
        simple.set("x", "0")
        simple.set("y", "0")
        floating.append(simple)
        for axis, offset in (
            ("H", block.bbox.x0),
            ("V", block.bbox.top),
        ):
            position = OxmlElement(f"wp:position{axis}")
            position.set("relativeFrom", "page")
            node = OxmlElement("wp:posOffset")
            node.text = str(round(offset * 12700))
            position.append(node)
            floating.append(position)
        for tag in ("wp:extent", "wp:effectExtent"):
            child = inline.find(qn(tag))
            if child is not None:
                floating.append(child)
        floating.append(OxmlElement("wp:wrapNone"))
        for tag in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
            child = inline.find(qn(tag))
            if child is not None:
                floating.append(child)
        inline.getparent().replace(inline, floating)

    @staticmethod
    def _float_table(table, block: TableBlock) -> None:
        properties = table._tbl.tblPr
        positioning = OxmlElement("w:tblpPr")
        positioning.set(qn("w:leftFromText"), "0")
        positioning.set(qn("w:rightFromText"), "0")
        positioning.set(qn("w:topFromText"), "0")
        positioning.set(qn("w:bottomFromText"), "0")
        positioning.set(qn("w:vertAnchor"), "page")
        positioning.set(qn("w:horzAnchor"), "page")
        positioning.set(qn("w:tblpX"), str(round(block.bbox.x0 * 20)))
        positioning.set(qn("w:tblpY"), str(round(block.bbox.top * 20)))
        properties.insert(0, positioning)

    @staticmethod
    def _layout_bands(blocks: Sequence[Block]) -> List[List[Block]]:
        bands: List[List[Block]] = []
        for block in sorted(blocks, key=lambda item: (item.bbox.top, item.bbox.x0)):
            if not bands:
                bands.append([block])
                continue
            band = bands[-1]
            band_top = min(item.bbox.top for item in band)
            band_bottom = max(item.bbox.bottom for item in band)
            if block.bbox.top <= band_bottom + 2 and block.bbox.bottom >= band_top - 2:
                band.append(block)
            else:
                bands.append([block])
        return bands

    @staticmethod
    def _layout_columns(band: Sequence[Block]) -> List[List[Block]]:
        columns: List[List[Block]] = []
        for block in sorted(band, key=lambda item: item.bbox.x0):
            if not columns:
                columns.append([block])
                continue
            right = max(item.bbox.x1 for item in columns[-1])
            if block.bbox.x0 <= right + 8:
                columns[-1].append(block)
            else:
                columns.append([block])
        return columns

    @staticmethod
    def _add_vertical_gap(document: Document, gap: float) -> None:
        # Only keep intentional whitespace. Small gaps are already covered by
        # Word's natural line/paragraph metrics and must not be double-counted.
        if gap <= 14:
            return
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(min(72, gap - 10))
        paragraph.paragraph_format.line_spacing = Pt(1)
        paragraph.add_run(" ").font.size = Pt(1)

    def _add_layout_row(
        self,
        document: Document,
        columns: Sequence[Sequence[Block]],
        page: PageIR,
    ) -> None:
        section = document.sections[-1]
        left_edge = section.left_margin.pt
        right_edge = page.width - section.right_margin.pt
        first_x = min(block.bbox.x0 for block in columns[0])
        last_x = max(block.bbox.x1 for block in columns[-1])
        spans_page = first_x < page.width * 0.35 and last_x > page.width * 0.65
        table_left = left_edge if spans_page else max(left_edge, first_x - 4)
        table_right = right_edge if spans_page else min(right_edge, last_x + 12)
        boundaries = [table_left]
        for left_column, right_column in zip(columns, columns[1:]):
            left_column_right = max(block.bbox.x1 for block in left_column)
            right_column_left = min(block.bbox.x0 for block in right_column)
            boundaries.append((left_column_right + right_column_left) / 2)
        boundaries.append(table_right)
        widths = [
            max(24.0, boundaries[index + 1] - boundaries[index])
            for index in range(len(columns))
        ]

        table = document.add_table(rows=1, cols=len(columns))
        table.alignment = (
            WD_TABLE_ALIGNMENT.RIGHT
            if table_left >= page.width * 0.5
            else WD_TABLE_ALIGNMENT.LEFT
        )
        table.autofit = False
        self._set_table_widths(table, widths)
        self._remove_table_borders(table)
        for cell, width, column in zip(table.rows[0].cells, widths, columns):
            cell.width = Pt(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins([cell], 0)
            initial = cell.paragraphs[0]
            for block in sorted(column, key=lambda item: item.bbox.top):
                self._add_block(cell, block, page, in_layout=True)
            if len(cell.paragraphs) > 1 and not initial.text:
                initial._element.getparent().remove(initial._element)

    def _add_block(
        self,
        container,
        block: Block,
        page: PageIR,
        in_layout: bool = False,
    ) -> None:
        if isinstance(block, ParagraphBlock):
            self._add_paragraph(container, block, page, in_layout)
        elif isinstance(block, TableBlock):
            section = container.sections[-1] if hasattr(container, "sections") else None
            self._add_table(container, block, section)
        elif isinstance(block, ImageBlock):
            self._add_image(container, block, page)

    def _add_paragraph(
        self,
        document,
        block: ParagraphBlock,
        page: PageIR,
        in_layout: bool = False,
    ) -> None:
        style = (
            f"Heading {block.heading_level}"
            if block.kind == "heading" and block.heading_level
            else ("List Number" if block.kind == "list" and block.list_ordered else "List Bullet")
            if block.kind == "list"
            else "Normal"
        )
        paragraph = document.add_paragraph(style=style)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        # Preserve baseline spacing while allowing Word to expand for tall glyphs.
        tops = sorted({round(span.bbox.top, 1) for span in block.spans if span.text.strip()})
        size = max((span.font_size for span in block.spans), default=11)
        leading = [b - a for a, b in zip(tops, tops[1:]) if size * .6 < b - a < size * 2]
        paragraph.paragraph_format.line_spacing = Pt(median(leading) if leading else size * 1.15)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        paragraph.paragraph_format.widow_control = False
        paragraph.alignment = {
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
            "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        }.get(block.alignment, WD_ALIGN_PARAGRAPH.LEFT)
        if block.kind != "list" and not in_layout:
            # Section margins already encode the page's left edge. Avoid stacking
            # another large indent that squeezes the right side of long lines.
            paragraph.paragraph_format.left_indent = Pt(0)
        else:
            paragraph.paragraph_format.left_indent = Inches(0.25 * block.list_level)

        spans = list(block.spans)
        for index, span in enumerate(spans):
            if not span.text:
                continue
            # Flowing pages should reflow; keep soft breaks only for stacked labels.
            text = span.text if in_layout or block.kind != "paragraph" else span.text.replace("\n", " ")
            if block.kind == "list" and index == 0:
                text = LIST_PREFIX.sub("", text, count=1)
            run = paragraph.add_run(text)
            # A small tracking tolerance accommodates PDF glyph positioning
            # without reducing the source font size or fixing line breaks.
            spacing = OxmlElement("w:spacing")
            spacing.set(qn("w:val"), "-2")
            run._r.get_or_add_rPr().append(spacing)
            run.bold = span.bold
            run.italic = span.italic
            run.font.name = self._safe_font(span.font_name)
            run.font.size = Pt(self._word_font_size(span.font_name, span.font_size))
            if span.color and not all(component > 245 for component in span.color):
                run.font.color.rgb = RGBColor(*span.color)

    @staticmethod
    def _add_table(document, block: TableBlock, section=None):
        columns = max((len(row) for row in block.rows), default=1)
        table = document.add_table(rows=len(block.rows), cols=columns)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        left = section.left_margin.pt if section is not None else 24.0
        content_width = (section.page_width.pt - left - section.right_margin.pt) if section is not None else 547.0
        offset = max(0, block.bbox.x0 - left)
        available = max(24.0, min(content_width - offset, block.bbox.width))
        widths = (
            [
                available * ratio
                for ratio in (0.40, 0.14, 0.08, 0.13, 0.155, 0.095)
            ]
            if columns == 6
            else [available * 0.44]
            + [available * 0.56 / (columns - 1)] * (columns - 1)
            if columns >= 4
            else [available / columns] * columns
        )
        DocxBuilder._set_table_widths(table, widths)
        preferred_width = table._tbl.tblPr.first_child_found_in("w:tblW")
        preferred_width.set(qn("w:type"), "auto")
        preferred_width.set(qn("w:w"), "0")
        indent = OxmlElement("w:tblInd")
        indent.set(qn("w:type"), "dxa")
        indent.set(qn("w:w"), str(round(offset * 20)))
        table._tbl.tblPr.append(indent)
        for row_index, values in enumerate(block.rows):
            for column_index in range(columns):
                value = values[column_index] if column_index < len(values) else ""
                cell = table.cell(row_index, column_index)
                cell.width = Pt(widths[column_index])
                cell.text = value
                set_cell_margins([cell], 55)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT
                    if column_index == 0
                    else WD_ALIGN_PARAGRAPH.CENTER
                )
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8.5)
                    run.bold = row_index == 0
                    if row_index == 0:
                        run.font.color.rgb = RGBColor(255, 255, 255)
                if row_index == 0:
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:fill"), "0F766E")
                    cell._tc.get_or_add_tcPr().append(shading)
        return table

    @staticmethod
    def _remove_table_borders(table) -> None:
        properties = table._tbl.tblPr
        borders = properties.first_child_found_in("w:tblBorders")
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            properties.append(borders)
        for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
            node = OxmlElement(f"w:{edge}")
            node.set(qn("w:val"), "nil")
            borders.append(node)

    @staticmethod
    def _set_table_widths(table, widths: Sequence[float]) -> None:
        table_width = table._tbl.tblPr.first_child_found_in("w:tblW")
        if table_width is None:
            table_width = OxmlElement("w:tblW")
            table._tbl.tblPr.append(table_width)
        table_width.set(qn("w:type"), "dxa")
        table_width.set(qn("w:w"), str(round(sum(widths) * 20)))

        for grid_column, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
            grid_column.set(qn("w:w"), str(round(width * 20)))
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Pt(width)
                cell_width = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
                if cell_width is None:
                    cell_width = OxmlElement("w:tcW")
                    cell._tc.get_or_add_tcPr().append(cell_width)
                cell_width.set(qn("w:type"), "dxa")
                cell_width.set(qn("w:w"), str(round(width * 20)))

    @staticmethod
    def _add_image(document: Document, block: ImageBlock, page: PageIR) -> None:
        if not block.path.exists():
            return
        section = document.sections[-1] if hasattr(document, "sections") else None
        available = section.page_width.pt - section.left_margin.pt - section.right_margin.pt if section is not None else page.width - 72
        width = min(max(block.bbox.width, 36), available)
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1
        paragraph.add_run().add_picture(
            str(block.path), width=Pt(width),
            height=Pt(block.bbox.height * width / max(1, block.bbox.width)),
        )

    @staticmethod
    def _safe_font(value: str) -> str:
        cleaned = re.sub(r"^[A-Z]{6}\+", "", value or "")
        cleaned = re.sub(r"[^A-Za-z0-9 ._-]", "", cleaned).strip()
        family = re.sub(
            r"(?i)[ _-]?(?:regular|roman|bolditalic|boldoblique|semibold|"
            r"bold|black|heavy|italic|oblique)$",
            "",
            cleaned,
        ).strip()
        aliases = {
            # Carlito is shipped privately by LibreOffice, so Word may replace
            # it with Times New Roman. Calibri is its Office counterpart.
            "carlito": "Calibri",
            "liberationsans": "Arial",
            "liberationserif": "Times New Roman",
        }
        return aliases.get(family.casefold(), family[:80]) or "Arial"

    @staticmethod
    def _word_font_size(font_name: str, font_size: float) -> float:
        # Keep the source size; global shrinking makes reports hard to read.
        return max(7.0, min(36.0, round(font_size * 2) / 2))


def set_cell_margins(cells: Iterable, margin_twips: int = 80) -> None:
    """OOXML helper kept public for advanced table reconstruction."""
    for cell in cells:
        properties = cell._tc.get_or_add_tcPr()
        margins = properties.first_child_found_in("w:tcMar")
        if margins is None:
            margins = OxmlElement("w:tcMar")
            properties.append(margins)
        for edge in ("top", "start", "bottom", "end"):
            node = OxmlElement(f"w:{edge}")
            node.set(qn("w:w"), str(margin_twips))
            node.set(qn("w:type"), "dxa")
            margins.append(node)
